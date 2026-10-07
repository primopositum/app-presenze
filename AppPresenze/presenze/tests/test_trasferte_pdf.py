from decimal import Decimal

from django.test import SimpleTestCase
from docx import Document
from docx.shared import Pt

from presenze.TrasfertePDF import (
    MONEY_CELL_FONT_SIZE_PT,
    _format_money,
    _shrink_long_money_cell_text,
)


class TestTrasfertePdfMoneyCells(SimpleTestCase):
    def test_format_money_uses_a_non_breaking_space_and_euro_symbol(self):
        self.assertEqual(_format_money(Decimal("1234.56")), "1234,56\u00a0\u20ac")

    def test_long_money_value_uses_a_smaller_font(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        run = cell.paragraphs[0].add_run("1.234,56\u00a0\u20ac")
        run.font.size = Pt(11)

        _shrink_long_money_cell_text(cell)

        self.assertEqual(run.font.size.pt, MONEY_CELL_FONT_SIZE_PT)

    def test_short_money_value_keeps_the_template_font_size(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        run = cell.paragraphs[0].add_run("12,34\u00a0\u20ac")
        run.font.size = Pt(11)

        _shrink_long_money_cell_text(cell)

        self.assertEqual(run.font.size.pt, 11)
